"""Figure-first literature reading and reporting workflow."""
from __future__ import annotations

from datetime import datetime
from html import escape
from pathlib import Path
from typing import Callable
from uuid import uuid4

from fastapi import APIRouter, Form, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse

from db import connect, now_iso
from runtime_paths import STORAGE_ROOT
from web.runtime import NOTE_IMAGE_DIR, UPLOAD_DIR

MAX_FIGURES = 12
FIGURE_ROOT = NOTE_IMAGE_DIR / "literature_figures"
PPT_ROOT = STORAGE_ROOT / "deliveries" / "literature_presentations"


def _clean(value: str, limit: int) -> str:
    return value.strip()[:limit]


def _fields(**values: str) -> dict[str, str]:
    limits = {
        "title": 240, "source": 600, "context": 1600, "question": 1600,
        "method": 1800, "key_figure": 1600, "evidence": 2400, "claim": 1800,
        "limits": 1800, "connection": 1800, "next_action": 800,
    }
    return {key: _clean(values.get(key, ""), limit) for key, limit in limits.items()}


def _report(fields: dict[str, str]) -> dict[str, object]:
    title = fields["title"] or "未命名论文"
    question = fields["question"] or "作者没有明确写出待验证问题"
    claim = fields["claim"] or "尚未填写结论"
    evidence = fields["evidence"] or "尚未定位关键图或数据"
    method = fields["method"] or "尚未填写方法与比较对象"
    limits = fields["limits"] or "尚未填写边界或替代解释"
    connection = fields["connection"] or "尚未连接到自己的课题"
    next_action = fields["next_action"] or "回到原文，补齐一张关键图的图注与数值。"
    slides = [
        ("1. 为什么值得听", fields["context"] or f"本篇要回答：{question}"),
        ("2. 作者的问题", question), ("3. 怎么做", method),
        ("4. 最关键的一张图", fields["key_figure"] or evidence),
        ("5. 证据是否够", evidence), ("6. 作者真正得到的结论", claim),
        ("7. 不应说得太满的地方", limits),
        ("8. 对我课题的下一步", connection + "\n下一动作：" + next_action),
    ]
    script = (
        f"我今天汇报《{title}》。它要解决的问题是：{question}。"
        f"作者用{method}来回答。最该看的证据是{evidence}。"
        f"因此我认为这篇论文支持的结论是：{claim}。"
        f"但它的边界是：{limits}。对我的课题，最有用的启发是：{connection}。"
        f"所以我的下一步是：{next_action}。"
    )
    return {"slides": slides, "script": script}


def _paper_choices() -> list[dict[str, object]]:
    with connect() as conn:
        return [
            dict(row) for row in conn.execute(
                """SELECT id,title,original_name,source FROM entries
                   WHERE status='active' AND file_path IS NOT NULL
                   AND (mime_type='application/pdf' OR lower(original_name) LIKE '%.pdf')
                   ORDER BY updated_at DESC LIMIT 80"""
            )
        ]


def _saved_reports() -> list[dict[str, object]]:
    with connect() as conn:
        return [
            dict(row) for row in conn.execute(
                """SELECT id,title,summary,content,updated_at FROM entries
                   WHERE kind='note' AND tags LIKE '%文献汇报%'
                   ORDER BY updated_at DESC LIMIT 12"""
            )
        ]


def _paper_file(entry_id: int) -> tuple[Path, dict[str, object]]:
    with connect() as conn:
        row = conn.execute(
            "SELECT id,title,file_path,original_name,source FROM entries WHERE id=? AND status='active'",
            (entry_id,),
        ).fetchone()
    if not row or not row["file_path"]:
        raise HTTPException(status_code=404, detail="未找到已归档的 PDF。")
    path = (UPLOAD_DIR / str(row["file_path"])).resolve()
    if UPLOAD_DIR.resolve() not in path.parents or not path.is_file() or path.suffix.lower() != ".pdf":
        raise HTTPException(status_code=404, detail="PDF 文件已移动或不可用。")
    return path, dict(row)


def _extract_figures(entry_id: int) -> list[str]:
    """Extract embedded raster figures; render pages as a useful fallback for vector plots."""
    pdf_path, _ = _paper_file(entry_id)
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise RuntimeError("当前构建缺少 PyMuPDF，无法从 PDF 提取图片。") from exc

    target = FIGURE_ROOT / f"entry_{entry_id}"
    target.mkdir(parents=True, exist_ok=True)
    for old in target.glob("*"):
        if old.is_file():
            old.unlink()
    urls: list[str] = []
    document = fitz.open(pdf_path)
    try:
        seen_xrefs: set[int] = set()
        for page_index, page in enumerate(document):
            for image_index, meta in enumerate(page.get_images(full=True), start=1):
                if len(urls) >= MAX_FIGURES:
                    break
                xref = int(meta[0])
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                payload = document.extract_image(xref)
                data = payload.get("image", b"")
                if len(data) < 12_000:  # ignore icons, rules and tiny logos
                    continue
                ext = str(payload.get("ext", "png")).lower()
                if ext not in {"png", "jpg", "jpeg", "webp"}:
                    ext = "png"
                name = f"embedded_p{page_index + 1}_{image_index}.{ext}"
                (target / name).write_bytes(data)
                urls.append(f"/media/note-images/literature_figures/entry_{entry_id}/{name}")
            if len(urls) >= MAX_FIGURES:
                break

        # Many science figures are vector drawings. Render figure-labelled pages so
        # the user always gets a slide-ready candidate rather than an empty gallery.
        if len(urls) < 4:
            for page_index, page in enumerate(document):
                text = page.get_text("text")[:700].lower()
                if "fig." not in text and "figure" not in text:
                    continue
                name = f"page_p{page_index + 1}.png"
                page.get_pixmap(matrix=fitz.Matrix(1.55, 1.55), alpha=False).save(target / name)
                urls.append(f"/media/note-images/literature_figures/entry_{entry_id}/{name}")
                if len(urls) >= MAX_FIGURES:
                    break
    finally:
        document.close()
    return urls


def _safe_figure_paths(entry_id: int, figure_urls: list[str]) -> list[tuple[str, Path]]:
    prefix = f"/media/note-images/literature_figures/entry_{entry_id}/"
    root = (FIGURE_ROOT / f"entry_{entry_id}").resolve()
    result: list[tuple[str, Path]] = []
    seen: set[str] = set()
    for url in figure_urls:
        if not url.startswith(prefix) or url in seen:
            continue
        name = url.removeprefix(prefix)
        path = (root / name).resolve()
        if root not in path.parents or not path.is_file():
            continue
        seen.add(url)
        result.append((url, path))
    return result[:8]


def _available_figures(entry_id: int) -> list[str]:
    root = FIGURE_ROOT / f"entry_{entry_id}"
    if not entry_id or not root.is_dir():
        return []
    return [
        f"/media/note-images/literature_figures/entry_{entry_id}/{path.name}"
        for path in sorted(root.iterdir()) if path.is_file()
    ][:MAX_FIGURES]


def _template(request: Request, context: Callable[..., dict], *, fields: dict[str, str], report=None, figures=None, paper_id=0, selected_figures=None):
    return context(
        request, "literature_report", report=report, fields=fields,
        saved_reports=_saved_reports(), paper_choices=_paper_choices(),
        figures=figures or [], selected_paper_id=paper_id,
        selected_figure_urls=set(selected_figures or []),
    )


def _make_pptx(title: str, report: dict[str, object], figures: list[tuple[str, Path]]) -> Path:
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    PPT_ROOT.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    def add_text(slide, text: str, x: float, y: float, w: float, h: float, size: int, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(44, 54, 62)
        return box

    title_slide = presentation.slides.add_slide(blank)
    add_text(title_slide, title or "文献汇报", 0.9, 1.2, 11.4, 1.1, 30, True)
    add_text(title_slide, "图像驱动汇报 · 问道科研", 0.9, 2.5, 8.5, .5, 15)
    add_text(title_slide, str(report["script"])[:420], 0.9, 3.35, 10.9, 1.5, 17)

    for index, (_url, image_path) in enumerate(figures, start=1):
        slide = presentation.slides.add_slide(blank)
        add_text(slide, f"关键图 {index}：先讲它证明什么", .65, .35, 11.9, .5, 24, True)
        with Image.open(image_path) as image:
            ratio = image.width / max(image.height, 1)
        max_w, max_h = 9.1, 5.65
        width = min(max_w, max_h * ratio)
        height = width / ratio
        left = .65 + (max_w - width) / 2
        top = 1.15 + (max_h - height) / 2
        slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))
        note = str(report["slides"][3][1]) if index == 1 else "这张图的：对象 → 对照 → 趋势 → 可支持的结论。"
        box = add_text(slide, note, 10.0, 1.4, 2.65, 4.55, 16)
        box.text_frame.word_wrap = True
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    closing = presentation.slides.add_slide(blank)
    add_text(closing, "结论、边界与自己的下一步", .8, .55, 11.4, .7, 27, True)
    for index, (heading, content) in enumerate(report["slides"][5:], start=0):
        add_text(closing, str(heading), 1.0, 1.65 + index * 1.65, 3.1, .4, 18, True)
        add_text(closing, str(content), 1.0, 2.12 + index * 1.65, 10.8, 1.0, 16)
    output = PPT_ROOT / f"文献汇报_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:6]}.pptx"
    presentation.save(output)
    return output


def register_literature_report_routes(app, templates, context: Callable[..., dict], flash):
    router = APIRouter()

    @router.get("/literature-report", response_class=HTMLResponse, name="literature_report_page")
    def literature_report_page(request: Request, paper_id: int = 0):
        return templates.TemplateResponse(
            request=request, name="literature_report.html",
            context=_template(request, context, fields={}, paper_id=paper_id),
        )

    @router.post("/literature-report/extract", response_class=HTMLResponse, name="literature_report_extract")
    def literature_report_extract(
        request: Request, paper_id: int = Form(0), title: str = Form(""), source: str = Form(""),
        context_text: str = Form(""), question: str = Form(""), method: str = Form(""),
        key_figure: str = Form(""), evidence: str = Form(""), claim: str = Form(""),
        limits: str = Form(""), connection: str = Form(""), next_action: str = Form(""),
    ):
        fields = _fields(title=title, source=source, context=context_text, question=question, method=method,
                         key_figure=key_figure, evidence=evidence, claim=claim, limits=limits,
                         connection=connection, next_action=next_action)
        if not paper_id:
            flash(request, "先从知识库选择一份 PDF，再提取关键图。", "error")
            figures = []
        else:
            try:
                _path, paper = _paper_file(paper_id)
                fields["title"] = fields["title"] or str(paper["title"])
                fields["source"] = fields["source"] or str(paper["source"] or paper["original_name"] or "")
                figures = _extract_figures(paper_id)
                flash(request, f"已提取 {len(figures)} 张候选图。优先选真正支撑结论的 1–3 张。")
            except (OSError, RuntimeError) as exc:
                figures = []
                flash(request, f"图片提取失败：{exc}", "error")
        return templates.TemplateResponse(
            request=request, name="literature_report.html",
            context=_template(request, context, fields=fields, figures=figures, paper_id=paper_id),
        )

    @router.post("/literature-report/draft", response_class=HTMLResponse, name="literature_report_draft")
    def literature_report_draft(
        request: Request, paper_id: int = Form(0), figure_urls: list[str] = Form(default=[]),
        title: str = Form(""), source: str = Form(""), context_text: str = Form(""), question: str = Form(""),
        method: str = Form(""), key_figure: str = Form(""), evidence: str = Form(""), claim: str = Form(""),
        limits: str = Form(""), connection: str = Form(""), next_action: str = Form(""),
    ):
        fields = _fields(title=title, source=source, context=context_text, question=question, method=method,
                         key_figure=key_figure, evidence=evidence, claim=claim, limits=limits,
                         connection=connection, next_action=next_action)
        figures = _safe_figure_paths(paper_id, figure_urls) if paper_id else []
        return templates.TemplateResponse(
            request=request, name="literature_report.html",
            context=_template(request, context, fields=fields, report=_report(fields),
                              figures=_available_figures(paper_id), paper_id=paper_id,
                              selected_figures=[url for url, _path in figures]),
        )

    @router.post("/literature-report/save", name="literature_report_save")
    def literature_report_save(
        request: Request, paper_id: int = Form(0), figure_urls: list[str] = Form(default=[]),
        title: str = Form(""), source: str = Form(""), context_text: str = Form(""), question: str = Form(""),
        method: str = Form(""), key_figure: str = Form(""), evidence: str = Form(""), claim: str = Form(""),
        limits: str = Form(""), connection: str = Form(""), next_action: str = Form(""),
    ):
        fields = _fields(title=title, source=source, context=context_text, question=question, method=method,
                         key_figure=key_figure, evidence=evidence, claim=claim, limits=limits,
                         connection=connection, next_action=next_action)
        report = _report(fields)
        selected = _safe_figure_paths(paper_id, figure_urls) if paper_id else []
        body = "".join(f"<h2>{escape(str(heading))}</h2><p>{escape(str(content)).replace(chr(10), '<br>')}</p>" for heading, content in report["slides"])
        if selected:
            body += "<h2>汇报关键图</h2>" + "".join(
                f'<p><img src="{escape(url, quote=True)}" alt="文献汇报关键图"></p>' for url, _path in selected
            )
        body += f"<h2>3 分钟口述稿</h2><p>{escape(str(report['script']))}</p>"
        timestamp = now_iso()
        with connect() as conn:
            cursor = conn.execute(
                """INSERT INTO entries(title,kind,domain,tags,summary,content,source,extract_status,content_format,created_at,updated_at)
                   VALUES (?,'note','文献汇报','文献汇报,证据卡,可编辑',?,?,?,'ready','html',?,?)""",
                (fields["title"] or "未命名文献汇报", fields["claim"][:280], body, fields["source"], timestamp, timestamp),
            )
            entry_id = int(cursor.lastrowid)
            conn.commit()
        flash(request, "汇报卡已归档为可编辑知识卡；结论变了就回来修改，而不是重复建一份。")
        return RedirectResponse(request.url_for("entry_view", entry_id=entry_id), status_code=303)

    @router.post("/literature-report/pptx", name="literature_report_pptx")
    def literature_report_pptx(
        request: Request, paper_id: int = Form(0), figure_urls: list[str] = Form(default=[]),
        title: str = Form(""), source: str = Form(""), context_text: str = Form(""), question: str = Form(""),
        method: str = Form(""), key_figure: str = Form(""), evidence: str = Form(""), claim: str = Form(""),
        limits: str = Form(""), connection: str = Form(""), next_action: str = Form(""),
    ):
        fields = _fields(title=title, source=source, context=context_text, question=question, method=method,
                         key_figure=key_figure, evidence=evidence, claim=claim, limits=limits,
                         connection=connection, next_action=next_action)
        selected = _safe_figure_paths(paper_id, figure_urls) if paper_id else []
        if not selected:
            flash(request, "请至少选一张已提取的关键图，再生成 PPT。", "error")
            return RedirectResponse(request.url_for("literature_report_page", paper_id=paper_id), status_code=303)
        output = _make_pptx(fields["title"], _report(fields), selected)
        return FileResponse(output, filename=output.name, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")

    app.include_router(router)
